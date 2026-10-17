"""Validate the hand-labeled workbook and export evals/golden_50_labeled.csv (no model calls).

Usage: python scripts/export_golden.py
Fails loudly on blank or invalid labels and on evidence quotes that are not exact substrings.
"""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipeline.common import INTENTS, TOPICS, csv_rows  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "labeling/golden_50_labeling.xlsx"
SRC = ROOT / "data/raw/golden_50_to_label.csv"
OUT = ROOT / "evals/golden_50_labeled.csv"
FIELDS = ["review_id", "topic", "intent", "severity", "sentiment", "evidence_quote", "entities",
          "needs_review", "ambiguous", "notes"]


def as_bool(v):
    return str(v).strip().upper() == "TRUE"


def main():
    source = {r["review_id"]: r["review_text"] for r in csv_rows(SRC)}
    ws = load_workbook(XLSX, data_only=True)["Label here"]
    header = [c.value for c in ws[1]]
    col = {name: header.index(name) for name in FIELDS}
    out, problems = [], []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        rid = row[col["review_id"]]
        if not rid:
            continue
        get = lambda k: "" if row[col[k]] is None else str(row[col[k]]).strip()  # noqa: E731
        rec = {k: get(k) for k in FIELDS}
        if rid not in source:
            problems.append(f"row {n}: unknown review_id")
        if rec["topic"] not in TOPICS:
            problems.append(f"row {n}: topic '{rec['topic']}' missing/invalid")
        if rec["intent"] not in INTENTS:
            problems.append(f"row {n}: intent '{rec['intent']}' missing/invalid")
        if rec["severity"] not in {"1", "2", "3", "4", "5"}:
            problems.append(f"row {n}: severity '{rec['severity']}' missing/invalid")
        try:
            if not -1 <= float(rec["sentiment"]) <= 1:
                raise ValueError
        except ValueError:
            problems.append(f"row {n}: sentiment '{rec['sentiment']}' missing/invalid")
        q = row[col["evidence_quote"]] or ""
        if q and q not in source.get(rid, ""):
            problems.append(f"row {n}: evidence_quote is not an exact substring of the review")
        rec["evidence_quote"] = q
        rec["needs_review"] = str(as_bool(rec["needs_review"])).lower()
        rec["ambiguous"] = str(as_bool(rec["ambiguous"])).lower()
        out.append(rec)
    if len(out) != 50:
        problems.append(f"expected 50 rows, found {len(out)}")
    if problems:
        print("Not exported. Fix these in the workbook:\n  " + "\n  ".join(problems))
        sys.exit(1)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(out)
    print(f"Exported {len(out)} human labels to {OUT}")


if __name__ == "__main__":
    main()
