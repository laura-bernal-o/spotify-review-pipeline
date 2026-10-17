"""Build the human labeling workbook for the golden 50 (no model calls).

Usage: python scripts/make_golden_sheet.py
Output: labeling/golden_50_labeling.xlsx
After labeling, run scripts/export_golden.py to write evals/golden_50_labeled.csv.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipeline.common import INTENTS, TOPICS, csv_rows  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data/raw/golden_50_to_label.csv"
OUT = ROOT / "labeling/golden_50_labeling.xlsx"

FONT = "Arial"
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # yellow = you fill in
BODY = Font(name=FONT, size=11)
GREY = Font(name=FONT, size=10, color="808080")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

COLUMNS = [  # header, width, input?
    ("#", 5, False), ("review_id", 14, False), ("review_text", 60, False),
    ("topic", 13, True), ("intent", 14, True), ("severity", 9, True), ("sentiment", 10, True),
    ("evidence_quote", 35, True), ("entities", 18, True), ("needs_review", 12, True),
    ("ambiguous", 11, True), ("notes", 30, True),
    ("stars (ignore while labeling)", 12, False),
]


def build():
    rows = list(csv_rows(SRC))
    assert len(rows) == 50, len(rows)
    wb = Workbook()

    ws = wb.active
    ws.title = "Label here"
    ws.append([c[0] for c in COLUMNS])
    for i, (name, width, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEAD_FILL, WRAP, THIN
        ws.column_dimensions[cell.column_letter].width = width
    for n, r in enumerate(rows, 1):
        ws.append([n, r["review_id"], r["review_text"], None, None, None, None, None, None, None, None, None,
                   r["review_rating"]])
        for i, (_, _, is_input) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=n + 1, column=i)
            cell.font, cell.alignment, cell.border = BODY, WRAP, THIN
            if is_input:
                cell.fill = INPUT_FILL
        ws.cell(row=n + 1, column=len(COLUMNS)).font = GREY
    ws.freeze_panes = "D2"
    last = len(rows) + 1

    def dropdown(col, values, prompt):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Not allowed", error="Pick from the list.",
                            promptTitle=prompt.split(":")[0], prompt=prompt, showInputMessage=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    dropdown("D", TOPICS, "topic: main product area (see Definitions tab)")
    dropdown("E", INTENTS, "intent: earliest in list wins if several apply")
    dropdown("F", ["1", "2", "3", "4", "5"], "severity: reported impact, not anger or stars")
    dropdown("G", ["-1", "-0.5", "0", "0.5", "1"], "sentiment: -1 very negative to 1 very positive")
    dropdown("J", ["FALSE", "TRUE"], "needs_review: TRUE if context is missing to judge")
    dropdown("K", ["FALSE", "TRUE"], "ambiguous: TRUE if two labels are equally defensible")

    d = wb.create_sheet("Definitions")
    d.column_dimensions["A"].width = 16
    d.column_dimensions["B"].width = 95
    lines = [
        ("HOW TO USE", None),
        ("Yellow cells", "Fill these in on the 'Label here' tab. Use the dropdowns. Leave other cells alone."),
        ("Read the text", "Label only from the review text. Do not look at stars. Do not use Claude/AI to label."),
        ("evidence_quote", "Copy-paste the exact words from the review that support your topic (short reviews: whole text)."),
        ("entities", "Optional: specific features/things named, comma-separated (e.g. 'shuffle, ads')."),
        ("notes", "Optional: why you chose a label when it was a close call."),
        ("", None),
        ("TOPIC", "Pick the most severe specific problem; if tied, the first one mentioned. "
                  "For praise, the first specific praised feature; general praise = other."),
        ("access", "Login, signup, password or account access"),
        ("usability", "Navigation, controls, layout, queue/playlist management, ad interruptions"),
        ("playback", "Playback failure, crashes, lag, connection failures, audio quality, resource use"),
        ("downloads", "Downloading, saved music, offline listening, disappearing downloads"),
        ("catalog", "Missing songs/artists, search/discovery, recommendations, lyrics availability"),
        ("billing", "Price, charges, subscriptions, paywalls, premium entitlement; explicitly premium-only controls. "
                    "Mentioning a paid plan alone is NOT billing."),
        ("support", "Contacting support and the support response"),
        ("other", "General praise/criticism, unrelated content, or no supported specific topic"),
        ("", None),
        ("INTENT", "If several apply, the first in this list wins."),
        ("cancellation", "Explicitly leaving, uninstalling, cancelling, or threatening to"),
        ("complaint", "Negative experience, including mixed praise/criticism. Generic 'bad app' = complaint."),
        ("request", "Wants a change, without reporting a failure"),
        ("praise", "Positive"),
        ("unclear", "Meaningless/unrelated text, or bare boycott slogans with no product complaint or personal departure"),
        ("", None),
        ("SEVERITY", "Stars, angry language, or cancellation intent alone do NOT raise severity."),
        ("1", "No reported problem: praise, neutral/unclear content, or a pure feature request"),
        ("2", "Dislike, generic criticism, minor annoyance, cosmetic issue; no functional loss"),
        ("3", "A degraded or restricted function; some use or workaround remains"),
        ("4", "A clearly blocked core task, e.g. cannot log in or cannot play music"),
        ("5", "Explicit serious financial, privacy, or data harm (an expensive plan or a crash alone is not enough)"),
        ("", None),
        ("SENTIMENT", "-1 very negative, -0.5 negative, 0 neutral/mixed, 0.5 positive, 1 very positive"),
        ("", None),
        ("EXAMPLES", "Made-up reviews (not from the dataset), showing the expected format:"),
        ("Example 1", "'Downloaded songs stop playing when I go offline.' -> downloads | complaint | 3 | -0.5 | "
                      "quote: whole text | needs_review FALSE"),
        ("Example 2", "'Can't log in since the update, I'm cancelling premium.' -> access | cancellation | 4 | -1 | "
                      "quote: 'Can't log in since the update' (cancelling does not raise severity; the blocked login does)"),
        ("Example 3", "'Love the Discover Weekly playlists!' -> catalog | praise | 1 | 1 | quote: whole text"),
        ("Example 4", "'Please add a sleep timer' -> usability | request | 1 | 0 | quote: whole text"),
        ("Example 5", "'Boycott Spotify!!!' -> other | unclear | 1 | -0.5 | needs_review TRUE"),
    ]
    for r, (a, b) in enumerate(lines, 1):
        d.cell(row=r, column=1, value=a).font = Font(name=FONT, bold=True)
        cb = d.cell(row=r, column=2, value=b)
        cb.font, cb.alignment = Font(name=FONT), WRAP
        if a in ("HOW TO USE", "TOPIC", "INTENT", "SEVERITY", "SENTIMENT", "EXAMPLES"):
            d.cell(row=r, column=1).fill = HEAD_FILL
            d.cell(row=r, column=1).font = HEAD
        if a == "Yellow cells":
            d.cell(row=r, column=1).fill = INPUT_FILL

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT} ({len(rows)} reviews)")


if __name__ == "__main__":
    build()
